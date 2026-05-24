from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
import pptx.oxml.ns as nsmap
from lxml import etree
import copy, os

IMGS = "/home/user/wendy/imgs"

# ── colours ──────────────────────────────────────────────
BG      = RGBColor(0xF7, 0xF4, 0xEF)
DARK    = RGBColor(0x1E, 0x1A, 0x17)
TERRA   = RGBColor(0xC4, 0x61, 0x2A)
CEL     = RGBColor(0x6B, 0x9E, 0x8E)
GOLD    = RGBColor(0xB0, 0x8A, 0x25)
GRAY    = RGBColor(0x8A, 0x85, 0x7E)
LGRAY   = RGBColor(0xD8, 0xD3, 0xCC)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
INK     = RGBColor(0x2C, 0x18, 0x10)

W = Inches(13.33)   # 16:9 width
H = Inches(7.5)

prs = Presentation()
prs.slide_width  = W
prs.slide_height = H

blank_layout = prs.slide_layouts[6]  # completely blank

# ── helpers ──────────────────────────────────────────────
def slide():
    s = prs.slides.add_slide(blank_layout)
    fill_slide(s, BG)
    return s

def fill_slide(s, color):
    bg = s.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def box(s, x, y, w, h, color):
    shp = s.shapes.add_shape(1, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    return shp

def txt(s, text, x, y, w, h,
        size=18, bold=False, color=DARK,
        align=PP_ALIGN.LEFT, wrap=True, italic=False):
    txb = s.shapes.add_textbox(x, y, w, h)
    tf  = txb.text_frame
    tf.word_wrap = wrap
    p   = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = 'PingFang SC'
    return txb

def rule(s, x, y, w, color=TERRA, h_pt=2):
    ln = s.shapes.add_shape(1, x, y, w, Pt(h_pt))
    ln.fill.solid()
    ln.fill.fore_color.rgb = color
    ln.line.fill.background()

def chip(s, text, x, y):
    bx = s.shapes.add_shape(1, x, y, Inches(1.1), Inches(0.3))
    bx.fill.background()
    bx.line.color.rgb = LGRAY
    bx.line.width = Pt(0.8)
    tf = bx.text_frame
    tf.word_wrap = False
    p  = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.size = Pt(9)
    run.font.color.rgb = GRAY
    run.font.name = 'PingFang SC'

# ══════════════════════════════════════════════════════════
# SLIDE 1 · COVER
# ══════════════════════════════════════════════════════════
s1 = slide()

# right panel warm tint
box(s1, Inches(8.5), 0, Inches(4.83), H, RGBColor(0xF0,0xEB,0xE4))

# decorative concentric circles (SVG-like using stacked rings)
cx, cy = Inches(11.5), Inches(4.2)
for r, op in [(3.2,0.06),(2.5,0.09),(1.8,0.13),(1.2,0.18),(0.6,0.25)]:
    ri = Inches(r)
    shp = s1.shapes.add_shape(9, cx-ri, cy-ri, ri*2, ri*2)  # 9=oval
    shp.fill.background()
    shp.line.color.rgb = TERRA
    shp.line.width = Pt(1.0)

# pottery vessel silhouette (simplified rectangle stand-in)
box(s1, Inches(10.6), Inches(2.5), Inches(0.9), Inches(2.8),
    RGBColor(0xC4,0x61,0x2A)).__class__  # just set opacity via XML later

# left text
txt(s1, "WAYFINDING SYSTEM SURVEY · XI'AN · 2025",
    Inches(0.7), Inches(0.7), Inches(7), Inches(0.4),
    size=8, color=TERRA)

txt(s1, "西安", Inches(0.7), Inches(1.3), Inches(6), Inches(1.1),
    size=60, bold=False, color=DARK)
txt(s1, "公共文化机构", Inches(0.7), Inches(2.2), Inches(6), Inches(0.9),
    size=44, bold=False, color=DARK)
txt(s1, "导视系统", Inches(0.7), Inches(3.0), Inches(6), Inches(0.9),
    size=50, bold=True, color=TERRA)

txt(s1, "实地考察与调研报告",
    Inches(0.7), Inches(4.0), Inches(6), Inches(0.5),
    size=16, color=GRAY)

rule(s1, Inches(0.7), Inches(5.0), Inches(0.4))

chip(s1, "博物馆", Inches(0.7), Inches(5.25))
chip(s1, "文化馆", Inches(1.9), Inches(5.25))
chip(s1, "图书馆", Inches(3.1), Inches(5.25))

txt(s1, "陶瓷体验博物馆导视创新研究",
    Inches(0.7), Inches(6.6), Inches(6), Inches(0.4),
    size=10, color=LGRAY)

# ══════════════════════════════════════════════════════════
# SLIDE 2 · KEYWORDS
# ══════════════════════════════════════════════════════════
s2 = slide()
fill_slide(s2, WHITE)

# giant bg character
txt(s2, "导", Inches(5.5), Inches(-1.2), Inches(8), Inches(8),
    size=320, bold=True, color=RGBColor(0xF0,0xED,0xE8))

txt(s2, "WAYFINDING · KEYWORDS",
    Inches(0.8), Inches(0.7), Inches(8), Inches(0.4), size=9, color=TERRA)
txt(s2, "导视", Inches(0.8), Inches(1.1), Inches(4), Inches(0.9),
    size=52, bold=False, color=DARK)
txt(s2, "关键词", Inches(2.5), Inches(1.1), Inches(4), Inches(0.9),
    size=52, bold=True, color=TERRA)

# Tier labels + words
tiers = [
    ("核心概念", [("引导",38,True,DARK),("叙事",38,True,TERRA),
                  ("识别",38,True,DARK),("体验",38,True,CEL)]),
    ("设计原则", [("系统性",22,False,TERRA),("一致性",22,False,TERRA),
                  ("在地性",22,False,TERRA),("可读性",22,False,TERRA),("无障碍",22,False,TERRA)]),
    ("功能维度", [("空间定向",15,False,GRAY),("文化表达",15,False,GRAY),
                  ("信息层级",15,False,GRAY),("情绪引导",15,False,GRAY),("品牌塑造",15,False,GRAY)]),
    ("物质语言", [("材质语言",11,False,LGRAY),("色彩编码",11,False,LGRAY),
                  ("形式语言",11,False,LGRAY),("比例关系",11,False,LGRAY),
                  ("字体系统",11,False,LGRAY),("图形符号",11,False,LGRAY)]),
]

y_starts = [2.3, 3.35, 4.2, 4.9]
for i,(label, words) in enumerate(tiers):
    txt(s2, label, Inches(0.8), Inches(y_starts[i]), Inches(1.1),
        Inches(0.5), size=8, color=LGRAY)
    xoff = Inches(1.95)
    for (w, sz, bd, col) in words:
        t = txt(s2, w, xoff, Inches(y_starts[i]-0.05), Inches(1.6),
                Inches(0.6), size=sz, bold=bd, color=col)
        xoff += Pt(sz)*len(w)*1.05 + Inches(0.18)

# thin horizontal divider after tier 1
rule(s2, Inches(0.8), Inches(3.25), Inches(11.5), LGRAY, 1)

# ══════════════════════════════════════════════════════════
# SLIDE 3 · MUSEUM CASES
# ══════════════════════════════════════════════════════════
s3 = slide()

# left dark panel
box(s3, 0, 0, Inches(2.3), H, DARK)
txt(s3, "MUSEUM\nWAYFINDING", Inches(0.15), Inches(0.6), Inches(2),
    Inches(0.7), size=7, color=RGBColor(0x60,0x60,0x60))
txt(s3, "中国\n博物馆\n导视内容", Inches(0.18), Inches(1.1), Inches(1.95),
    Inches(2.2), size=28, bold=False, color=WHITE)
txt(s3, "导视内容包括：", Inches(0.18), Inches(3.5), Inches(1.95),
    Inches(0.4), size=8, color=RGBColor(0x60,0x60,0x60))
items = ["导览总图","楼层索引","展厅标牌","方向箭头",
         "展品说明","服务设施","禁止标识","安全出口"]
for i,it in enumerate(items):
    txt(s3, it, Inches(0.18), Inches(3.95+i*0.33), Inches(1.95),
        Inches(0.3), size=9, color=RGBColor(0x80,0x80,0x80))

museums = [
    ("故宫博物院", "Beijing · Imperial Palace",
     [RGBColor(0xC4,0x1E,0x3A),RGBColor(0xB8,0x86,0x0B),
      RGBColor(0x2C,0x18,0x10),RGBColor(0xF0,0xE6,0x8C)],
     "云纹卷草 · 宫廷红金",
     "朱红、金色为主色系，融入卷云纹、回纹等传统纹样；\n字体选用宋体+隶书，铜质鎏金材质，彰显皇家礼制氛围。",
     "gugong"),
    ("苏州博物馆", "Suzhou · I.M.Pei Design",
     [RGBColor(0x1A,0x1A,0x1A),RGBColor(0xFF,0xFF,0xFF),
      RGBColor(0x88,0x88,0x88),RGBColor(0xD4,0xC5,0xA9)],
     "几何抽象 · 黑白灰",
     "贝聿铭设计语言延伸，菱形几何为核心图形；\n极简黑白色系呼应粉墙黛瓦，字体纤细，大量留白。",
     "suzhou"),
    ("上海博物馆", "Shanghai · Bronze Culture",
     [RGBColor(0x5C,0x3D,0x2E),RGBColor(0x7B,0x9E,0x87),
      RGBColor(0xB0,0x8A,0x25),RGBColor(0xF5,0xF0,0xE8)],
     "青铜纹饰 · 棕金体系",
     "以青铜器轮廓为视觉标识原型，提取饕餮纹、云雷纹再设计；\n铜绿色与暖棕色形成稳重历史感，兼顾现代简洁性。",
     "shanghai"),
    ("良渚博物院", "Hangzhou · Jade Culture",
     [RGBColor(0x5B,0x8A,0x7A),RGBColor(0x8F,0xAF,0x9F),
      RGBColor(0xC4,0xB8,0x9A),RGBColor(0xF8,0xF4,0xEE)],
     "玉琮形态 · 青绿玉色",
     "从玉琮（方圆组合）提炼核心图形语言，象征天地融合；\n玉绿色系传递史前文明的神秘与沉静，字体轻盈现代。",
     "liangzhu"),
]

card_h = Inches(1.62)
for i, (name, loc, palette, tags, desc, img_key) in enumerate(museums):
    y = Inches(0.18) + i * (card_h + Inches(0.06))
    # card bg
    box(s3, Inches(2.4), y, Inches(10.7), card_h, WHITE)
    # accent bar
    box(s3, Inches(2.4), y, Inches(0.08), card_h, palette[0])
    # photo on right
    img_path = f"{IMGS}/{img_key}.jpg"
    if os.path.exists(img_path):
        s3.shapes.add_picture(img_path, Inches(10.18), y+Inches(0.08),
                              Inches(2.84), card_h-Inches(0.16))
    # name
    txt(s3, name, Inches(2.6), y+Inches(0.15), Inches(2),
        Inches(0.45), size=17, bold=True, color=DARK)
    txt(s3, loc, Inches(2.6), y+Inches(0.55), Inches(2),
        Inches(0.3), size=8, color=GRAY)
    # palette swatches
    for j, col in enumerate(palette):
        bx = s3.shapes.add_shape(9, Inches(2.6)+j*Inches(0.28),
                                  y+Inches(0.95), Inches(0.22), Inches(0.22))
        bx.fill.solid()
        bx.fill.fore_color.rgb = col
        bx.line.fill.background()
    # tags
    txt(s3, tags, Inches(4.8), y+Inches(0.15), Inches(2.5),
        Inches(0.35), size=10, color=TERRA)
    # desc (narrowed to leave room for photo)
    txt(s3, desc, Inches(4.8), y+Inches(0.5), Inches(5.2),
        Inches(0.95), size=10.5, color=GRAY)

# ══════════════════════════════════════════════════════════
# SLIDE 4 · THEME CONNECTION
# ══════════════════════════════════════════════════════════
s4 = slide()
fill_slide(s4, WHITE)

# left terracotta panel
box(s4, 0, 0, Inches(2.8), H, TERRA)
txt(s4, "04", Inches(0.15), Inches(0.3), Inches(2.5), Inches(1.8),
    size=100, bold=True, color=RGBColor(0xFF,0xFF,0xFF))
# set opacity manually is complex; just use lighter shade
txt(s4, "导视与\n博物馆主题\n的联系",
    Inches(0.18), Inches(1.8), Inches(2.45), Inches(2.2),
    size=26, color=WHITE)
txt(s4, "Theme · Connection",
    Inches(0.18), Inches(4.2), Inches(2.5), Inches(0.4),
    size=9, color=RGBColor(0xFF,0xCC,0xAA))
txt(s4, "博物馆的文化属性\n决定导视的\n色彩、纹样与材质",
    Inches(0.18), Inches(4.8), Inches(2.5), Inches(1.5),
    size=11, color=RGBColor(0xFF,0xCC,0xAA))

# table header
headers = ["馆藏类型","色彩语言","图形/形式语言","代表案例"]
hx = [3.0, 4.5, 6.8, 10.2]
for hd, x in zip(headers, hx):
    txt(s4, hd, Inches(x), Inches(0.5), Inches(2.5),
        Inches(0.35), size=9, color=GRAY)
rule(s4, Inches(3.0), Inches(0.9), Inches(10.1), DARK, 1.5)

rows = [
    ("遗址考古",
     [RGBColor(0xA0,0x52,0x2D),RGBColor(0x6B,0x8E,0x7A),RGBColor(0x8B,0x73,0x55)],
     "土层纹理·考古断面\n地层线条·出土器型",
     "秦始皇帝陵博物院\n金沙遗址博物馆"),
    ("宫廷历史",
     [RGBColor(0xC4,0x1E,0x3A),RGBColor(0xB8,0x86,0x0B),RGBColor(0x2C,0x18,0x10)],
     "云纹·龙凤·回纹\n铜质鎏金·宫廷礼仪",
     "故宫博物院\n沈阳故宫博物院"),
    ("自然科技",
     [RGBColor(0x3A,0x7C,0xA5),RGBColor(0x52,0xB7,0x88),RGBColor(0xE8,0xF4,0xF8)],
     "有机曲线·数据可视化\n网格系统·生物形态",
     "中国科学技术馆\n自然博物馆"),
    ("艺术当代",
     [RGBColor(0xF5,0xF5,0xF5),RGBColor(0x1A,0x1A,0x1A),RGBColor(0xE8,0x44,0x1A)],
     "极简几何·无衬线字体\n大留白·高纯度点缀",
     "中国美术馆\n苏州博物馆"),
    ("陶瓷工艺",
     [RGBColor(0xC4,0x61,0x2A),RGBColor(0x6B,0x9E,0x8E),RGBColor(0xB0,0x8A,0x25)],
     "器型轮廓·纹样提取\n窑变色彩·手工质感",
     "景德镇陶瓷博物馆\n▶ 你的陶瓷体验馆"),
]

for i,(rtype, colors, lang, eg) in enumerate(rows):
    y = Inches(1.1) + i*Inches(1.15)
    txt(s4, rtype, Inches(3.0), y, Inches(1.3), Inches(0.5),
        size=15, bold=True, color=TERRA if i==4 else DARK)
    # color dots
    for j,c in enumerate(colors):
        shp = s4.shapes.add_shape(9, Inches(4.5)+j*Inches(0.34), y+Inches(0.05),
                                   Inches(0.27), Inches(0.27))
        shp.fill.solid(); shp.fill.fore_color.rgb = c
        shp.line.fill.background()
    txt(s4, lang, Inches(6.8), y, Inches(3.0), Inches(0.55),
        size=10.5, color=TERRA if i==4 else GRAY)
    txt(s4, eg, Inches(10.2), y, Inches(2.8), Inches(0.55),
        size=10, color=TERRA if i==4 else RGBColor(0xB0,0x8A,0x25))
    if i < 4:
        rule(s4, Inches(3.0), y+Inches(0.72), Inches(10.1), LGRAY, 0.8)

# ══════════════════════════════════════════════════════════
# SLIDE 5 · FUNCTIONS
# ══════════════════════════════════════════════════════════
s5 = slide()

# header bar
box(s5, 0, 0, W, Inches(1.55), BG)
rule(s5, 0, Inches(1.55), W, LGRAY, 1)
txt(s5, "WAYFINDING · FUNCTIONS",
    Inches(0.7), Inches(0.35), Inches(6), Inches(0.4), size=9, color=TERRA)
txt(s5, "导视系统的", Inches(0.7), Inches(0.75), Inches(4), Inches(0.7),
    size=42, bold=False, color=DARK)
txt(s5, "作用", Inches(3.15), Inches(0.75), Inches(2), Inches(0.7),
    size=42, bold=True, color=DARK)

funcs = [
    ("01","空间定向",
     "帮助访客在陌生空间中快速建立方位感，明确[我在哪里/我要去哪里]，降低迷失感与焦虑。"),
    ("02","信息传递",
     "传达展览内容、开放时间、服务设施等实用信息，是文字说明与图形符号的综合信息界面。"),
    ("03","品牌塑造",
     "导视系统是机构视觉形象的物理延伸，通过统一的色彩、字体、材质强化文化品牌认知。"),
    ("04","情绪引导",
     "通过色调、材质与空间节奏，在参观动线上营造期待、沉浸、惊喜等情绪序列。"),
    ("05","安全疏导",
     "规范标注安全出口、紧急集合点、消防设施，确保紧急情况下人员安全疏散。"),
]
col_w = Inches(2.55)
for i,(num,name,desc) in enumerate(funcs):
    x = Inches(0.12) + i*col_w
    if i > 0:
        rule(s5, x-Inches(0.06), Inches(1.6), Inches(0.01), LGRAY, 200)
        ln = s5.shapes.add_shape(1, x-Inches(0.06), Inches(1.6),
                                  Pt(1), Inches(5.9))
        ln.fill.solid(); ln.fill.fore_color.rgb = LGRAY
        ln.line.fill.background()
    txt(s5, num, x+Inches(0.2), Inches(1.8), Inches(1.2), Inches(0.9),
        size=44, bold=True, color=LGRAY)
    txt(s5, name, x+Inches(0.2), Inches(2.9), Inches(2.2), Inches(0.5),
        size=20, bold=True, color=DARK)
    txt(s5, desc, x+Inches(0.2), Inches(3.5), Inches(2.2), Inches(2.5),
        size=11, color=GRAY)

# ══════════════════════════════════════════════════════════
# SLIDE 6 · CONSIDERATIONS
# ══════════════════════════════════════════════════════════
s6 = slide()
fill_slide(s6, WHITE)

# right dark panel
box(s6, Inches(9.5), 0, Inches(3.83), H, DARK)
txt(s6, "好的导视，\n让访客感觉\n不到它的存在",
    Inches(9.65), Inches(1.2), Inches(3.5), Inches(2.5),
    size=22, color=WHITE)
checks = [
    "不喧宾夺主，与展品相辅相成",
    "第一次来的人能自然找到路",
    "紧急情况下指引清晰无歧义",
    "承载文化身份，值得被拍照",
    "五年后仍与整体环境协调一致",
]
for i,c in enumerate(checks):
    txt(s6, "— "+c, Inches(9.65), Inches(4.0)+i*Inches(0.48),
        Inches(3.5), Inches(0.4), size=11, color=RGBColor(0xA0,0x98,0x90))

txt(s6, "DESIGN CONSIDERATIONS",
    Inches(0.7), Inches(0.5), Inches(6), Inches(0.4), size=9, color=TERRA)
txt(s6, "导视系统", Inches(0.7), Inches(0.9), Inches(5), Inches(0.7),
    size=40, bold=False, color=DARK)
txt(s6, "注意事项", Inches(3.55), Inches(0.9), Inches(3), Inches(0.7),
    size=40, bold=True, color=DARK)

considers = [
    ("01","视觉层级与可读性",
     "主/副标题/说明三级层级清晰。中文正文不低于14px，关键信息对比度≥4.5:1（WCAG AA）。"),
    ("02","无障碍设计",
     "考虑色盲用户（避免纯红绿区分）、视障人士（盲文铭牌）、轮椅视角（标牌高1.0–1.6m）。"),
    ("03","多语言与国际化",
     "提供中英双语，注意英文版面更长，排版需预留弹性空间。"),
    ("04","材质与空间协调",
     "材质应与建筑对话。户外需防UV防腐；室内可选亚克力、铝板或木质材料。"),
    ("05","系统一致性",
     "全馆遵循统一设计规范（色值、字体、图标、尺寸），避免各楼层各自为政。"),
    ("06","更新与维护弹性",
     "展览定期更换，标识系统需支持局部更新，避免每次改展都全面重制。"),
]
positions = [(0.65,1.85),(5.0,1.85),(0.65,3.7),(5.0,3.7),(0.65,5.5),(5.0,5.5)]
for (cx,cy),(num,title,desc) in zip(positions, considers):
    bx = s6.shapes.add_shape(1, Inches(cx), Inches(cy), Inches(4.1), Inches(1.55))
    bx.fill.background()
    bx.line.color.rgb = LGRAY
    bx.line.width = Pt(0.8)
    txt(s6, num+" ·", Inches(cx+0.15), Inches(cy+0.12), Inches(1), Inches(0.3),
        size=9, color=TERRA)
    txt(s6, title, Inches(cx+0.15), Inches(cy+0.38), Inches(3.8), Inches(0.38),
        size=15, bold=True, color=DARK)
    txt(s6, desc, Inches(cx+0.15), Inches(cy+0.75), Inches(3.8), Inches(0.7),
        size=10, color=GRAY)

# ══════════════════════════════════════════════════════════
# SLIDE 7 · XI'AN SURVEY
# ══════════════════════════════════════════════════════════
s7 = slide()

# header
box(s7, 0, 0, W, Inches(1.3), BG)
rule(s7, 0, Inches(1.3), W, LGRAY, 1)
txt(s7, "FIELD RESEARCH · XI'AN",
    Inches(0.7), Inches(0.22), Inches(6), Inches(0.35), size=9, color=TERRA)
txt(s7, "西安", Inches(0.7), Inches(0.55), Inches(2), Inches(0.65),
    size=38, bold=False, color=DARK)
txt(s7, "实地考察", Inches(1.62), Inches(0.55), Inches(4), Inches(0.65),
    size=38, bold=True, color=TERRA)
txt(s7, "6处文化机构 · 导视系统观察记录",
    Inches(9.0), Inches(0.85), Inches(4.0), Inches(0.35),
    size=10, color=GRAY, align=PP_ALIGN.RIGHT)

venues = [
    ("01","陕西历史博物馆","综合历史 · 唐风建筑",
     RGBColor(0xB8,0x86,0x0B),
     ["暖土黄+金属金，呼应唐代宫廷气韵",
      "楼层索引牌融入唐草纹装饰边框",
      "户外立式导览牌采用仿唐碑形制"],
     ["唐风纹样","暖金色系","铜质材料"], "shaanxi"),
    ("02","秦始皇帝陵博物院","遗址考古 · 室外为主",
     RGBColor(0x7B,0x4F,0x3A),
     ["赭石+铜绿，呼应兵马俑本体色调",
      "大尺寸图形标识适配户外宽阔空间",
      "兵马俑剪影作为主要图形元素"],
     ["户外大尺度","多语言","图形化"], "terracotta"),
    ("03","西安博物院","城市历史 · 小雁塔景区",
     RGBColor(0x4A,0x67,0x41),
     ["小雁塔轮廓作为标识原型图形",
      "现代简约风，中性石灰色为主",
      "室内外导视系统风格统一度高"],
     ["建筑轮廓","现代简约","双功能区"], "xian_bwy"),
    ("04","碑林博物馆","书法碑刻 · 文字文化",
     RGBColor(0x2C,0x2C,0x2C),
     ["以石刻拓印质感为背景纹理",
      "字体选用隶书与楷书",
      "墨黑+石灰色调，沉静典雅"],
     ["书法字体","拓印质感","墨色系"], "beilin"),
    ("05","西安市图书馆","公共图书馆 · 功能导向",
     RGBColor(0x2B,0x6C,0xB0),
     ["色彩编码系统：不同阅览区用色区分",
      "功能性优先，导视清晰直接",
      "文化属性体现不足，风格较通用"],
     ["色彩分区","功能导向","现代系统"], "library"),
    ("06","西安非遗文化馆","非物质文化遗产展示",
     RGBColor(0xC4,0x1E,0x3A),
     ["剪纸、皮影等民俗纹样入标",
      "色彩饱和度高，节日感浓郁",
      "互动导视（扫码查看制作视频）"],
     ["民俗纹样","高饱和色","互动导视"], "noncultural"),
]

col_w2 = Inches(4.3)
for i,(num,name,vtype,col,findings,tags,img_key) in enumerate(venues):
    row = i // 3
    col_i = i % 3
    x = Inches(0.06) + col_i*(col_w2+Inches(0.1))
    y = Inches(1.38) + row*Inches(2.98)
    # card bg
    card = s7.shapes.add_shape(1, x, y, col_w2, Inches(2.88))
    card.fill.background()
    card.line.color.rgb = LGRAY
    card.line.width = Pt(0.8)
    # photo strip across top of card
    img_path = f"{IMGS}/{img_key}.jpg"
    if os.path.exists(img_path):
        s7.shapes.add_picture(img_path, x+Inches(0.06), y+Inches(0.06),
                              col_w2-Inches(0.12), Inches(1.0))
    # number + color dot (overlaid on photo)
    txt(s7, num, x+Inches(0.12), y+Inches(0.08), Inches(0.55), Inches(0.45),
        size=22, bold=True, color=WHITE)
    dot = s7.shapes.add_shape(9, x+col_w2-Inches(0.5), y+Inches(0.12),
                               Inches(0.26), Inches(0.26))
    dot.fill.solid(); dot.fill.fore_color.rgb = col
    dot.line.fill.background()
    # name & type below photo
    txt(s7, name, x+Inches(0.18), y+Inches(1.1), col_w2-Inches(0.36),
        Inches(0.42), size=15, bold=True, color=DARK)
    txt(s7, vtype, x+Inches(0.18), y+Inches(1.5), col_w2-Inches(0.36),
        Inches(0.25), size=8, color=GRAY)
    for j,f in enumerate(findings):
        txt(s7, "— "+f, x+Inches(0.18), y+Inches(1.78)+j*Inches(0.34),
            col_w2-Inches(0.36), Inches(0.3), size=9.5, color=GRAY)

# ══════════════════════════════════════════════════════════
# SLIDE 8 · CERAMIC INNOVATION
# ══════════════════════════════════════════════════════════
s8 = slide()
fill_slide(s8, WHITE)

# left dark panel
box(s8, 0, 0, Inches(3.2), H, INK)
txt(s8, "CERAMIC MUSEUM · INNOVATION",
    Inches(0.2), Inches(0.5), Inches(2.8), Inches(0.4),
    size=7, color=TERRA)
txt(s8, "陶瓷体验\n博物馆\n导视创新",
    Inches(0.2), Inches(1.0), Inches(2.8), Inches(2.4),
    size=30, color=WHITE)
txt(s8, "8个创新方向",
    Inches(0.2), Inches(6.8), Inches(2.5), Inches(0.4),
    size=10, color=RGBColor(0x60,0x60,0x60))

# vessel silhouette (simple shape stand-in)
vessel = s8.shapes.add_shape(9, Inches(0.7), Inches(3.2), Inches(1.8), Inches(2.8))
vessel.fill.solid(); vessel.fill.fore_color.rgb = TERRA
vessel.line.fill.background()
# Approximate the vessel with adjusted shape
vessel2 = s8.shapes.add_shape(9, Inches(0.85), Inches(3.1), Inches(1.5), Inches(0.5))
vessel2.fill.solid(); vessel2.fill.fore_color.rgb = INK
vessel2.line.fill.background()

txt(s8, "8 Innovation\nDirections",
    Inches(3.4), Inches(0.3), Inches(4), Inches(0.6),
    size=9, color=GRAY)

innovations = [
    ("01","纹样提取再设计",
     "从馆藏陶瓷（彩陶鱼纹、唐三彩花卉）提取图形，重新设计为导视图标系统。"),
    ("02","陶轮同心圆方向系统",
     "以陶轮旋转的同心圆为核心图形，通过圆弧断点位置暗示路径方向。"),
    ("03","窑变色彩分区编码",
     "将窑变釉色（天青、火红、铁锈棕、蜜黄）作为各功能区的专属色彩标识。"),
    ("04","制作体验流程导视",
     "将取土→揉泥→拉坯→修坯→上釉→入窑的工序设计为叙事性导视动线。"),
    ("05","手印 · 指纹标识",
     "以陶艺师手印为视觉原点，提炼指纹线条作为品牌图形，强调手作温度。"),
    ("06","陶瓷断面美学",
     "将陶器断面（胎土、化妆土、釉层）视觉化，形成独特分层纹理图形。"),
    ("07","开片裂纹装饰语言",
     "哥窑冰裂纹提炼为导视底纹与边框装饰，将缺陷美转化为辨识度符号。"),
    ("08","数字互动导视",
     "实体标牌嵌入NFC/AR码，扫码可查看制作视频、历史溯源与展品3D模型。"),
]

grid_x = [3.35, 6.55, 9.75, 12.95]
grid_y = [0.75, 4.1]
for i,(num,title,desc) in enumerate(innovations):
    row = i // 4
    col_i = i % 4
    x = Inches(grid_x[col_i] if col_i < len(grid_x) else grid_x[-1])
    y = Inches(grid_y[row])
    card = s8.shapes.add_shape(1, x, y, Inches(3.0), Inches(3.1))
    card.fill.solid(); card.fill.fore_color.rgb = BG
    card.line.fill.background()
    txt(s8, num+" ·", x+Inches(0.18), y+Inches(0.18), Inches(0.8),
        Inches(0.3), size=9, color=TERRA)
    txt(s8, title, x+Inches(0.18), y+Inches(0.52), Inches(2.65),
        Inches(0.55), size=14, bold=True, color=DARK)
    txt(s8, desc, x+Inches(0.18), y+Inches(1.1), Inches(2.65),
        Inches(1.8), size=10, color=GRAY)

# ══════════════════════════════════════════════════════════
prs.save('/home/user/wendy/西安文化机构导视系统调研报告.pptx')
print("Done!")
